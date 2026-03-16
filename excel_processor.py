"""
Procesor Excel pentru Campania de Email
Curs: Inteligenta Artificiala - De la Zero la Expert

Citeste un fisier Excel cu lista de contacte si genereaza
subiecte de email personalizate pentru fiecare destinatar.

Coloane asteptate in fisierul Excel:
  - prenume   (obligatoriu)
  - email     (obligatoriu)
  - ore       (optional, default: 24)
  - reducere  (optional, default: 40)
  - locuri    (optional, default: 7)
"""

import sys
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from email_subject_generator import generate_subjects


# ──────────────────────────────────────────────
# Citire fisier Excel
# ──────────────────────────────────────────────

def citeste_excel(cale_fisier: str) -> list[dict]:
    """
    Citeste contactele dintr-un fisier Excel.

    Returneaza o lista de dict-uri cu cheile:
      prenume, email, ore, reducere, locuri
    """
    if not os.path.exists(cale_fisier):
        raise FileNotFoundError(f"Fisierul '{cale_fisier}' nu a fost gasit.")

    wb = openpyxl.load_workbook(cale_fisier)
    ws = wb.active

    # Citeste antetul (primul rand)
    antete = []
    for cell in ws[1]:
        valoare = str(cell.value).strip().lower() if cell.value else ""
        antete.append(valoare)

    if "prenume" not in antete or "email" not in antete:
        raise ValueError(
            "Fisierul Excel trebuie sa contina coloanele 'prenume' si 'email'."
        )

    contacte = []
    for rand in ws.iter_rows(min_row=2, values_only=True):
        # Sari randurile goale
        if not any(rand):
            continue

        rand_dict = dict(zip(antete, rand))

        # Valori obligatorii
        prenume_raw = rand_dict.get("prenume")
        email_raw = rand_dict.get("email")
        prenume = str(prenume_raw).strip() if prenume_raw is not None else ""
        email = str(email_raw).strip() if email_raw is not None else ""
        if not prenume or not email:
            continue

        contacte.append({
            "prenume": prenume,
            "email": email,
            "ore": int(rand_dict.get("ore") or 24),
            "reducere": int(rand_dict.get("reducere") or 40),
            "locuri": int(rand_dict.get("locuri") or 7),
        })

    return contacte


# ──────────────────────────────────────────────
# Scriere rezultate in Excel
# ──────────────────────────────────────────────

CULORI_CATEGORII = {
    "Urgenta_Pura":          "FFD7E4",   # roz deschis
    "Curiozitate_plus_Urgenta": "D7E8FF",  # albastru deschis
    "Beneficiu_plus_Urgenta":   "D7FFE4",  # verde deschis
    "Empowerment_plus_Urgenta": "FFF3D7",  # galben deschis
}


def _stil_antet(cell, culoare_hex="2D6A9F"):
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor=culoare_hex)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _bordura(cell):
    linie = Side(style="thin", color="AAAAAA")
    cell.border = Border(left=linie, right=linie, top=linie, bottom=linie)


def salveaza_rezultate_excel(contacte: list[dict], cale_iesire: str):
    """
    Genereaza subiecte pentru fiecare contact si salveaza
    rezultatele intr-un fisier Excel formatat.
    """
    wb = openpyxl.Workbook()

    # ── Foaia 1: Toate subiectele ──────────────────────────────
    ws_toate = wb.active
    ws_toate.title = "Toate Subiectele"

    antet = ["Prenume", "Email", "Categorie", "Nr.", "Subiect Email"]
    for col, titlu in enumerate(antet, start=1):
        cell = ws_toate.cell(row=1, column=col, value=titlu)
        _stil_antet(cell)
        _bordura(cell)

    ws_toate.row_dimensions[1].height = 30
    ws_toate.column_dimensions["A"].width = 14
    ws_toate.column_dimensions["B"].width = 28
    ws_toate.column_dimensions["C"].width = 28
    ws_toate.column_dimensions["D"].width = 6
    ws_toate.column_dimensions["E"].width = 75

    rand_curent = 2
    for contact in contacte:
        subiecte = generate_subjects(
            contact["prenume"],
            contact["ore"],
            contact["reducere"],
            contact["locuri"],
        )
        nr = 1
        for categorie, lista_subiecte in subiecte.items():
            culoare = CULORI_CATEGORII.get(categorie, "FFFFFF")
            for subiect in lista_subiecte:
                date_rand = [
                    contact["prenume"],
                    contact["email"],
                    categorie.replace("_", " "),
                    nr,
                    subiect,
                ]
                for col, valoare in enumerate(date_rand, start=1):
                    cell = ws_toate.cell(row=rand_curent, column=col, value=valoare)
                    cell.fill = PatternFill("solid", fgColor=culoare)
                    cell.alignment = Alignment(vertical="center", wrap_text=(col == 5))
                    _bordura(cell)
                nr += 1
                rand_curent += 1

    # ── Foaia 2: Rezumat per contact ──────────────────────────
    ws_rezumat = wb.create_sheet("Rezumat Contacte")
    antet_rez = ["Prenume", "Email", "Ore", "Reducere %", "Locuri", "Total Subiecte"]
    for col, titlu in enumerate(antet_rez, start=1):
        cell = ws_rezumat.cell(row=1, column=col, value=titlu)
        _stil_antet(cell, culoare_hex="1E7A4F")
        _bordura(cell)

    ws_rezumat.row_dimensions[1].height = 30
    for i, latime in enumerate([14, 28, 8, 12, 8, 16], start=1):
        ws_rezumat.column_dimensions[get_column_letter(i)].width = latime

    for rand_idx, contact in enumerate(contacte, start=2):
        subiecte = generate_subjects(
            contact["prenume"],
            contact["ore"],
            contact["reducere"],
            contact["locuri"],
        )
        total = sum(len(v) for v in subiecte.values())
        date = [
            contact["prenume"],
            contact["email"],
            contact["ore"],
            contact["reducere"],
            contact["locuri"],
            total,
        ]
        culoare = "EAF4EE" if rand_idx % 2 == 0 else "FFFFFF"
        for col, valoare in enumerate(date, start=1):
            cell = ws_rezumat.cell(row=rand_idx, column=col, value=valoare)
            cell.fill = PatternFill("solid", fgColor=culoare)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            _bordura(cell)

    wb.save(cale_iesire)
    print(f"Rezultate salvate in: {cale_iesire}")


# ──────────────────────────────────────────────
# Creare fisier template Excel
# ──────────────────────────────────────────────

def creeaza_template(cale_fisier: str = "contacte_template.xlsx"):
    """Genereaza un fisier Excel template cu date demo."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacte"

    antete = ["prenume", "email", "ore", "reducere", "locuri"]
    for col, titlu in enumerate(antete, start=1):
        cell = ws.cell(row=1, column=col, value=titlu)
        _stil_antet(cell)
        _bordura(cell)

    ws.row_dimensions[1].height = 28
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 10

    date_demo = [
        ("Maria",    "maria@exemplu.ro",    24, 40, 7),
        ("Ana",      "ana@exemplu.ro",      48, 35, 5),
        ("Elena",    "elena@exemplu.ro",    12, 50, 3),
        ("Cristina", "cristina@exemplu.ro", 36, 40, 10),
        ("Ioana",    "ioana@exemplu.ro",    24, 45, 8),
    ]

    for rand_idx, (prenume, email, ore, reducere, locuri) in enumerate(date_demo, start=2):
        culoare = "F0F8FF" if rand_idx % 2 == 0 else "FFFFFF"
        for col, valoare in enumerate([prenume, email, ore, reducere, locuri], start=1):
            cell = ws.cell(row=rand_idx, column=col, value=valoare)
            cell.fill = PatternFill("solid", fgColor=culoare)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            _bordura(cell)

    # Nota explicativa
    ws.cell(row=8, column=1, value="Nota:").font = Font(bold=True, color="666666")
    ws.cell(row=9, column=1, value="ore     = ore ramase pana la expirarea ofertei (default: 24)")
    ws.cell(row=10, column=1, value="reducere = procentul de reducere (default: 40)")
    ws.cell(row=11, column=1, value="locuri  = locuri disponibile ramase (default: 7)")
    for r in range(9, 12):
        ws.cell(row=r, column=1).font = Font(italic=True, color="888888", size=9)
    ws.merge_cells("A9:E9")
    ws.merge_cells("A10:E10")
    ws.merge_cells("A11:E11")

    wb.save(cale_fisier)
    print(f"Template creat: {cale_fisier}")


# ──────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────

def main():
    # Daca este furnizat un fisier Excel ca argument, il proceseaza
    if len(sys.argv) > 1:
        fisier_intrare = sys.argv[1]
        fisier_iesire = sys.argv[2] if len(sys.argv) > 2 else "rezultate_subiecte.xlsx"

        print(f"Citesc contactele din: {fisier_intrare}")
        contacte = citeste_excel(fisier_intrare)
        print(f"  -> {len(contacte)} contacte gasite.")

        salveaza_rezultate_excel(contacte, fisier_iesire)
    else:
        # Fara argumente: creeaza template demo si il proceseaza
        print("Niciun fisier specificat. Se creeaza si proceseaza template-ul demo...\n")
        creeaza_template("contacte_template.xlsx")

        contacte = citeste_excel("contacte_template.xlsx")
        print(f"  -> {len(contacte)} contacte gasite.")
        salveaza_rezultate_excel(contacte, "rezultate_subiecte.xlsx")

        print("\nUtilizare cu propriul fisier Excel:")
        print("  python3 excel_processor.py <fisier_intrare.xlsx> [fisier_iesire.xlsx]")
        print("\nStructura asteptata a fisierului Excel:")
        print("  Coloane obligatorii: prenume, email")
        print("  Coloane optionale:   ore (default 24), reducere (default 40), locuri (default 7)")


if __name__ == "__main__":
    main()
